"""
build_testcase.py — Generator file test case (XLSX có styling) cho skill create-testcase-figma-prd-v2.
Xuất đúng cấu trúc + dropdown/màu Test Type, Platform, Priority, Test Result theo template:
assets/templates/testcase-template-v2.csv (template này cũng được sinh từ chính file này — xem export_template_csv).

Cột `Test Type` (Positive / Negative / Edge) ở vị trí C và cột `Platform` (Both / Android / iOS) ở vị trí D
→ mỗi test case là tuple 10 field. Case dùng chung 2 nền tảng ghi Both;
bước nào khác nhau giữa 2 OS thì ghi "Android: … / iOS: …" ngay trong bước, không tách case.

Cách dùng:
    from build_testcase import build_testcase_xlsx

    testcases = [
        # (Test Case #, Module/Group test, Test Type, Platform, Test Scenario, Test Cases Description,
        #  Pre-condition, Steps, Expected Result, Priority)
        ("TC_LOGIN_001", "Đăng nhập", "Positive", "Both", "Đăng nhập thành công",
         "Verify user đăng nhập thành công với tài khoản hợp lệ",
         "Tài khoản đã đăng ký", "1. Mở Login\n2. Nhập đúng thông tin\n3. Nhấn Đăng nhập",
         "Đăng nhập thành công, chuyển sang Home", "High"),
        ...
    ]

    build_testcase_xlsx(
        project_name="...", function_name="...",
        written_by="...", written_date="...", reviewed_by="...", reviewed_date="...",
        execute_by="...", execute_date="...", srs_link="...",
        testcases=testcases, output_path="testcase-login.xlsx",
    )

RTM (ma trận truy vết 3 cấp Feature → Flow → Rule): truyền thêm `rtm=[...]` (hoặc key "rtm" trong từng sheet)
→ mỗi sheet test case có thêm 1 sheet "RTM - <tên sheet>" ngay sau nó. Mỗi dòng RTM là tuple 10 field:
    (ID, Cấp, Thuộc, Nội dung, Nguồn, Flow / E2E, Positive, Negative, Edge, Ghi chú)
    - Cấp: Feature / Flow / Rule. Thuộc: ID dòng cha (Flow thuộc Feature, Rule thuộc Flow hoặc Feature).
    - Feature/Flow: điền mã case đi hết luồng vào "Flow / E2E", để trống P/N/E.
    - Rule: điền P/N/E (nhiều mã cách nhau dấu phẩy, hoặc "N/A — <lý do>"), để trống "Flow / E2E".
Mã case được đánh lại liên tục (001, 002…) theo thứ tự truyền vào, mọi chỗ nhắc mã cũ đổi theo
→ xoá / chèn case thoải mái, không cần tự sửa số.
Kiểm tra RTM sau khi build: python3 check_rtm.py <file>.xlsx [<file-qna>.md]

Sinh lại template CSV khi đổi cột:
    python build_testcase.py --template assets/templates/testcase-template-v2.csv
"""

import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

FONT = "Arial"

HEADER = [
    "Test Case #", "Module/Group test", "Test Type", "Platform", "Test Scenario", "Test Cases Description",
    "Pre-condition", "Steps", "Expected Result", "Priority", "Test Result", "Actual Result",
]
N_COLS = len(HEADER)


def _col(title):
    """Chữ cái cột Excel theo tên header — tránh hardcode 'H', 'I'... khi thêm/bớt cột."""
    return get_column_letter(HEADER.index(title) + 1)


COL_TC_ID = _col("Test Case #")
COL_TEST_TYPE = _col("Test Type")
COL_PLATFORM = _col("Platform")
COL_PRIORITY = _col("Priority")
COL_TEST_RESULT = _col("Test Result")
COL_ACTUAL = _col("Actual Result")

C_HEADER_BG = "1F4E79"
C_HEADER_TXT = "FFFFFF"
C_LABEL_BG = "EFEFEF"

# màu Test Type — tông riêng, không trùng màu Priority/Test Result để khỏi đọc nhầm nghĩa
TEST_TYPE_OPTIONS = ["Positive", "Negative", "Edge"]
TEST_TYPE_STYLE = {
    "Positive": {"bg": "DDEBF7", "txt": "1F4E79"},
    "Negative": {"bg": "FCE4D6", "txt": "843C0C"},
    "Edge": {"bg": "E4DFEC", "txt": "5B2C6F"},
}

# Platform — Both = 1 case dùng chung iOS + Android; Android / iOS = hành vi chỉ có trên 1 nền tảng
PLATFORM_OPTIONS = ["Both", "Android", "iOS"]
PLATFORM_STYLE = {
    "Both": {"bg": "F2F2F2", "txt": "404040"},
    "Android": {"bg": "E2F0D9", "txt": "38761D"},
    "iOS": {"bg": "DEEAF6", "txt": "1F3864"},
}

# màu Priority (áp dụng qua Conditional Formatting để tự đổi khi đổi giá trị dropdown)
PRIORITY_OPTIONS = ["High", "Medium", "Low"]
PRIORITY_STYLE = {
    "High": {"bg": "FFC7CE", "txt": "9C0006"},
    "Medium": {"bg": "FFEB9C", "txt": "9C6500"},
    "Low": {"bg": "C6EFCE", "txt": "006100"},
}

# màu Test Result (áp dụng qua Conditional Formatting để tự đổi khi tester chọn dropdown)
TEST_RESULT_OPTIONS = ["Passed", "Failed", "Blocked", "To-do"]
TEST_RESULT_STYLE = {
    "Passed": {"bg": "C6EFCE", "txt": "006100"},
    "Failed": {"bg": "FFC7CE", "txt": "9C0006"},
    "Blocked": {"bg": "FFEB9C", "txt": "9C6500"},
    "To-do": {"bg": "D9D9D9", "txt": "595959"},
}

# sheet RTM
RTM_HEADER = ["ID", "Cấp", "Thuộc", "Nội dung (trích spec/design)", "Nguồn",
              "Flow / E2E", "Positive", "Negative", "Edge", "Ghi chú"]
RTM_LEVELS = ["Feature", "Flow", "Rule"]
RTM_WIDTHS = [9, 10, 9, 48, 18, 18, 18, 22, 22, 26]
RTM_LEVEL_FILL = {"Feature": "BDD7EE", "Flow": "DDEBF7", "Rule": None}
RTM_LINK_LABEL = "Sheet test case"   # ô A1; B1 = tên sheet test case mà RTM này truy vết — check_rtm.py dựa vào đây
RTM_HEADER_ROW = 3

# các cột dropdown: căn giữa + bold, màu set qua Conditional Formatting
DROPDOWN_COLS = {
    COL_TEST_TYPE: (TEST_TYPE_OPTIONS, TEST_TYPE_STYLE),
    COL_PLATFORM: (PLATFORM_OPTIONS, PLATFORM_STYLE),
    COL_PRIORITY: (PRIORITY_OPTIONS, PRIORITY_STYLE),
    COL_TEST_RESULT: (TEST_RESULT_OPTIONS, TEST_RESULT_STYLE),
}

_thin = Side(style="thin", color="B7B7B7")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

WIDTH_BY_TITLE = {
    "Test Case #": 20, "Module/Group test": 20, "Test Type": 12, "Platform": 11, "Test Scenario": 26,
    "Test Cases Description": 34, "Pre-condition": 24, "Steps": 34, "Expected Result": 34,
    "Priority": 12, "Test Result": 14, "Actual Result": 34,  # Actual Result = Expected Result
}

META_HEADER_ROW = 9  # dòng header cột (sau 7 dòng meta + 1 dòng trống, giống template gốc)
FIRST_DATA_ROW = 10

# cột của khối meta: nhãn trái (B), nhãn giữa (F), nhãn SUM (trên cột Test Result), công thức SUM (trên cột Actual Result)
META_LABEL_COL = 2
META_MID_LABEL_COL = 6
SUM_LABEL_COL = HEADER.index("Test Result") + 1
SUM_VALUE_COL = HEADER.index("Actual Result") + 1


def _style(cell, *, bold=False, color="000000", fill=None, size=11,
           halign="left", valign="center", wrap=False, border=True):
    cell.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    if border:
        cell.border = BORDER


def _sanitize_sheet_name(name, used):
    invalid = set('[]:*?/\\')
    cleaned = "".join(c for c in (name or "Sheet") if c not in invalid).strip() or "Sheet"
    cleaned = cleaned[:31]
    base, i = cleaned, 2
    while cleaned in used:
        suffix = f" ({i})"
        cleaned = base[: 31 - len(suffix)] + suffix
        i += 1
    return cleaned


def _meta_rows(project_name, function_name, written_by, written_date, reviewed_by, reviewed_date,
               execute_by, execute_date, srs_link, sum_values):
    """8 dòng meta đầu sheet. sum_values = [passed, failed, todo, blocked, total] (công thức hoặc rỗng)."""
    left = [("Project Name", project_name), ("Function Name", function_name),
            ("Written By", written_by), ("Written Date", written_date),
            ("Execute by", execute_by), ("Link SRS", srs_link)]
    # cùng dòng với Written By / Written Date / Execute by
    mid = {3: ("Reviewed By", reviewed_by), 4: ("Reviewed Date", reviewed_date),
           5: ("Execute Date", execute_date)}
    sum_labels = ["Passed", "Failed", "To-do", "Blocked", "Total of TC"]

    rows = [[""] * N_COLS for _ in range(8)]
    rows[0][SUM_VALUE_COL - 1] = "SUM"
    for i, (label, val) in enumerate(left, start=1):
        rows[i][META_LABEL_COL - 1] = label
        rows[i][META_LABEL_COL] = val
        if i in mid:
            rows[i][META_MID_LABEL_COL - 1] = mid[i][0]
            rows[i][META_MID_LABEL_COL] = mid[i][1]
        if i <= len(sum_labels):
            rows[i][SUM_LABEL_COL - 1] = sum_labels[i - 1]
            rows[i][SUM_VALUE_COL - 1] = sum_values[i - 1]
    return rows


def _validate(testcases):
    for tc in testcases:
        if len(tc) != 10:
            raise ValueError(f"Test case phải có 10 field (Test Type ở vị trí 3, Platform ở vị trí 4), nhận {len(tc)}: {tc[:1]}")
        if tc[2] not in TEST_TYPE_OPTIONS:
            raise ValueError(f"{tc[0]}: Test Type '{tc[2]}' không hợp lệ, chỉ nhận {TEST_TYPE_OPTIONS}")
        if tc[3] not in PLATFORM_OPTIONS:
            raise ValueError(f"{tc[0]}: Platform '{tc[3]}' không hợp lệ, chỉ nhận {PLATFORM_OPTIONS}")
        if tc[9] not in PRIORITY_OPTIONS:
            raise ValueError(f"{tc[0]}: Priority '{tc[9]}' không hợp lệ, chỉ nhận {PRIORITY_OPTIONS}")


TC_NUM_RE = re.compile(r"^(.*_)(\d+)$")


def _renumber(sheets):
    """Đánh lại số case liên tục theo thứ tự xuất hiện (từng prefix, xuyên mọi sheet),
    rồi đổi mọi chỗ nhắc mã cũ (Pre-condition, Steps, Expected, ô RTM) sang mã mới — 1 lượt, không thay chồng."""
    mapping, counters = {}, {}
    for sheet in sheets:
        for tc in sheet.get("testcases") or []:
            if tc[0] in mapping:
                raise ValueError(f"Trùng mã case: {tc[0]}")
            m = TC_NUM_RE.match(tc[0])
            if not m:
                mapping[tc[0]] = tc[0]
                continue
            prefix, width = m.group(1), len(m.group(2))
            counters[prefix] = counters.get(prefix, 0) + 1
            mapping[tc[0]] = f"{prefix}{counters[prefix]:0{width}d}"
    if all(k == v for k, v in mapping.items()):
        return
    pattern = re.compile(r"(?<![\w])(" + "|".join(sorted(map(re.escape, mapping), key=len, reverse=True)) + r")(?![\w])")
    sub = lambda v: pattern.sub(lambda m: mapping[m.group(1)], v) if isinstance(v, str) else v
    for sheet in sheets:
        sheet["testcases"] = [tuple(sub(v) for v in tc) for tc in sheet.get("testcases") or []]
        if sheet.get("rtm"):
            sheet["rtm"] = [tuple(sub(v) for v in row) for row in sheet["rtm"]]
    changed = [f"{k}→{v}" for k, v in mapping.items() if k != v]
    print(f"Đánh lại số {len(changed)} case: {', '.join(changed[:10])}{' …' if len(changed) > 10 else ''}")


def _validate_rtm(rtm):
    ids = set()
    for row in rtm:
        if len(row) != len(RTM_HEADER):
            raise ValueError(f"Dòng RTM phải có {len(RTM_HEADER)} field {RTM_HEADER}, nhận {len(row)}: {row[:1]}")
        if row[0] in ids:
            raise ValueError(f"RTM trùng ID: {row[0]}")
        ids.add(row[0])
        if row[1] not in RTM_LEVELS:
            raise ValueError(f"{row[0]}: Cấp '{row[1]}' không hợp lệ, chỉ nhận {RTM_LEVELS}")


def _build_rtm_sheet(ws, test_sheet_title, rtm):
    _style(ws.cell(row=1, column=1, value=RTM_LINK_LABEL), bold=True, fill=C_LABEL_BG)
    _style(ws.cell(row=1, column=2, value=test_sheet_title), bold=True)
    for c, title in enumerate(RTM_HEADER, start=1):
        _style(ws.cell(row=RTM_HEADER_ROW, column=c, value=title),
               bold=True, color=C_HEADER_TXT, fill=C_HEADER_BG, halign="center", wrap=True)
    for i, row in enumerate(rtm):
        r = RTM_HEADER_ROW + 1 + i
        fill = RTM_LEVEL_FILL[row[1]]
        for c, val in enumerate(row, start=1):
            _style(ws.cell(row=r, column=c, value=val or None), bold=row[1] != "Rule" and c <= 4,
                   fill=fill, wrap=True, valign="top")
    for c, w in enumerate(RTM_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = f"A{RTM_HEADER_ROW + 1}"


def build_testcase_xlsx(project_name,
                         written_by, written_date, reviewed_by, reviewed_date,
                         execute_by, execute_date,
                         function_name=None, srs_link=None, testcases=None, rtm=None,
                         sheets=None, output_path="testcase.xlsx"):
    """Sinh file XLSX (có styling + dropdown) theo đúng format testcase-template-v2.csv.

    Dùng function_name/srs_link/testcases cho trường hợp 1 sheet duy nhất, hoặc dùng
    sheets=[{"sheet_name":, "function_name":, "srs_link":, "testcases":, "rtm":}, ...] khi cần
    nhiều sheet trong cùng 1 file (vd: 1 sheet / 1 file spec trong thư mục nhiều file).
    """
    if sheets is None:
        sheets = [{
            "sheet_name": function_name or "Test Case",
            "function_name": function_name,
            "srs_link": srs_link,
            "testcases": testcases or [],
            "rtm": rtm,
        }]
    sheets = [dict(s) for s in sheets]
    _renumber(sheets)
    for sheet in sheets:
        _validate(sheet.get("testcases") or [])
        if sheet.get("rtm"):
            _validate_rtm(sheet["rtm"])

    wb = Workbook()
    used_names = set()
    for idx, sheet in enumerate(sheets):
        name = _sanitize_sheet_name(sheet.get("sheet_name") or sheet.get("function_name"), used_names)
        used_names.add(name)
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = name
        _build_sheet(
            ws, project_name,
            sheet.get("function_name") or "",
            written_by, written_date, reviewed_by, reviewed_date,
            execute_by, execute_date,
            sheet.get("srs_link") or "",
            sheet.get("testcases") or [],
        )
        if sheet.get("rtm"):
            rtm_name = _sanitize_sheet_name(f"RTM - {name}", used_names)
            used_names.add(rtm_name)
            rtm_ws = wb.create_sheet()
            rtm_ws.title = rtm_name
            _build_rtm_sheet(rtm_ws, name, sheet["rtm"])

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _build_sheet(ws, project_name, function_name,
                  written_by, written_date, reviewed_by, reviewed_date,
                  execute_by, execute_date, srs_link, testcases):
    """Điền toàn bộ nội dung + styling của 1 sheet test case (meta rows, header, data, dropdown, CF)."""
    n = len(testcases)
    last_data_row = FIRST_DATA_ROW + max(n, 1) - 1

    def rng(col):
        return f"{col}{FIRST_DATA_ROW}:{col}{last_data_row}"

    tr = rng(COL_TEST_RESULT)
    sum_values = [
        f'=COUNTIF({tr},"Passed")',
        f'=COUNTIF({tr},"Failed")',
        # COUNTBLANK chỉ cộng khi có dữ liệu — sheet 0 case thì vùng rỗng 1 ô sẽ bị đếm nhầm thành 1 To-do
        f'=COUNTIF({tr},"To-do")+COUNTBLANK({tr})' if n else f'=COUNTIF({tr},"To-do")',
        f'=COUNTIF({tr},"Blocked")',
        f'=COUNTA({rng(COL_TC_ID)})',
    ]

    meta = _meta_rows(project_name, function_name, written_by, written_date, reviewed_by,
                      reviewed_date, execute_by, execute_date, srs_link, sum_values)
    for r, row_vals in enumerate(meta, start=1):
        for c, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=val if val != "" else None)
            is_label = c in (META_LABEL_COL, META_MID_LABEL_COL) and val not in ("", None)
            is_sum_label = c == SUM_LABEL_COL and val not in ("", None)
            is_sum_formula = c == SUM_VALUE_COL and isinstance(val, str) and val.startswith("=")
            if r == 1 and c == SUM_VALUE_COL:
                _style(cell, bold=True, color=C_HEADER_TXT, fill=C_HEADER_BG, halign="center")
            elif is_label or is_sum_label:
                _style(cell, bold=True, fill=C_LABEL_BG, halign="left")
            elif is_sum_formula:
                _style(cell, bold=True, fill=C_LABEL_BG, halign="center")
            else:
                _style(cell, halign="left")

    # header cột
    for c, title in enumerate(HEADER, start=1):
        _style(ws.cell(row=META_HEADER_ROW, column=c, value=title),
               bold=True, color=C_HEADER_TXT, fill=C_HEADER_BG, halign="center", wrap=True)

    # data rows — Test Result mặc định "To-do", Actual Result để trống cho tester tự ghi
    dropdown_idx = {HEADER.index(t) + 1 for t in ("Test Type", "Platform", "Priority", "Test Result")}
    for i, tc in enumerate(testcases):
        r = FIRST_DATA_ROW + i
        values = list(tc) + ["To-do", ""]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            # TẤT CẢ các cột đều wrap text
            if c in dropdown_idx:
                _style(cell, bold=True, halign="center", wrap=True, valign="center")
            else:
                _style(cell, halign="left", wrap=True, valign="top")
        # KHÔNG set ws.row_dimensions[r].height — để Excel/Google Sheets/LibreOffice tự auto-fit
        # chiều cao dòng khi mở file theo font render thật của máy người dùng.

    # dropdown (Data Validation) + Conditional Formatting cho các cột phân loại
    for col, (options, styles) in DROPDOWN_COLS.items():
        dv = DataValidation(
            type="list",
            formula1=f'"{",".join(options)}"',
            allow_blank=True,
            showDropDown=False,  # False = hiện mũi tên dropdown (quirk của openpyxl/Excel)
        )
        dv.error = f"Chỉ chọn 1 trong: {', '.join(options)}"
        dv.errorTitle = "Giá trị không hợp lệ"
        ws.add_data_validation(dv)
        dv.add(rng(col))

        for option, style in styles.items():
            ws.conditional_formatting.add(rng(col), FormulaRule(
                formula=[f'{col}{FIRST_DATA_ROW}="{option}"'],
                fill=PatternFill("solid", fgColor=style["bg"]),
                font=Font(name=FONT, bold=True, color=style["txt"]),
            ))

    for title, w in WIDTH_BY_TITLE.items():
        ws.column_dimensions[_col(title)].width = w

    # App 1 nền tảng: mọi case cùng 1 Platform → ẩn cột (vẫn giữ dữ liệu để import / automation đọc được)
    plat_i = HEADER.index("Platform")
    if testcases and len({tc[plat_i] for tc in testcases}) == 1:
        ws.column_dimensions[_col("Platform")].hidden = True

    _merge_scenarios(ws, testcases)
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"


def _merge_scenarios(ws, testcases):
    """Gộp ô Test Scenario cho các case liền nhau cùng Module + cùng tên Scenario (chỉ dòng đầu nhóm giữ tên —
    công cụ đọc file cần điền xuống tên nhóm cho các dòng sau)."""
    mod_i, sc_i = HEADER.index("Module/Group test"), HEADER.index("Test Scenario")
    col = sc_i + 1
    i = 0
    while i < len(testcases):
        key = (testcases[i][mod_i], testcases[i][sc_i])
        j = i
        while j + 1 < len(testcases) and (testcases[j + 1][mod_i], testcases[j + 1][sc_i]) == key:
            j += 1
        if j > i and key[1]:
            ws.merge_cells(start_row=FIRST_DATA_ROW + i, end_row=FIRST_DATA_ROW + j, start_column=col, end_column=col)
            ws.cell(FIRST_DATA_ROW + i, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        i = j + 1


def export_template_csv(output_path, blank_rows=88):
    """Sinh template CSV (BOM UTF-8) đúng cấu trúc file XLSX mà generator này xuất ra."""
    rows = _meta_rows("", "", "", "", "", "", "", "", "", [""] * 5)
    rows.append(list(HEADER))
    todo_idx = HEADER.index("Test Result")
    for _ in range(blank_rows):
        row = [""] * N_COLS
        row[todo_idx] = "To-do"
        rows.append(row)
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows(rows)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--template":
        print("template saved:", export_template_csv(sys.argv[2]))
        sys.exit(0)
    print("usage: python3 build_testcase.py --template <output.csv>  (dùng như thư viện: from build_testcase import build_testcase_xlsx)", file=sys.stderr)
    sys.exit(1)
