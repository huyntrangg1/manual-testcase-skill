"""
check_rtm.py — Kiểm tra RTM (ma trận truy vết 3 cấp Feature → Flow → Rule) đối chiếu với test case
trong cùng file XLSX do build_testcase.py sinh ra. Kết quả tất định, không phụ thuộc đánh giá của AI.

Cách dùng:
    python3 check_rtm.py <file>.xlsx [<file-qna>.md]

Kiểm tra cho từng cặp (sheet test case, sheet "RTM - ..."):
    1. Cấu trúc: ID không trùng, Cấp hợp lệ, "Thuộc" trỏ tới dòng có thật và đúng cấp
       (Flow thuộc Feature; Rule thuộc Flow hoặc Feature), mỗi Feature có ít nhất 1 Flow.
    2. Bao phủ: Feature/Flow có case ở cột "Flow / E2E"; Rule có đủ 3 ô Positive/Negative/Edge
       (mã case hoặc "N/A — <lý do>"). Ô Positive chỉ được N/A khi Nguồn bắt đầu bằng "Checklist"
       hoặc lý do ghi rõ "không có hành vi positive".
    3. Liên kết: mã case trong RTM có thật trong sheet test case; mọi case đều được RTM tham chiếu (không mồ côi).
    4. Test Type: case nằm ở ô Positive/Negative/Edge phải có đúng Test Type đó.
    5. Pre-condition (rule Pre-condition ở bước Viết case của SKILL.md): không trống; đúng 1 trong 3 dạng — danh sách bước "1. …",
       "Tiếp tục từ TC_xxx: …" (TC_xxx phải là case ngay trước), "Navigate tới màn …"; không chứa câu mơ hồ.
       Expected Result không trống; Steps và Expected đánh số giống hệt nhau (mỗi step 1 dòng Expected).
    6. Q&A (khi truyền file Q&A): mọi "(Q&A #n)" phải có câu #n trong file; Expected "Chờ confirm (Q&A #n)"
       mà câu #n đã "Đã chốt" → lỗi, viết lại Expected theo câu trả lời.

Mã thoát: 0 = đạt (0 lỗi), 1 = có lỗi, 2 = không đọc được file / không có sheet RTM.
"""

import re
import sys

from openpyxl import load_workbook

TC_HEADER_ROW = 9
RTM_LINK_LABEL = "Sheet test case"
RTM_HEADER_ROW = 3
RTM_COLS = ["ID", "Cấp", "Thuộc", "Nội dung", "Nguồn", "Flow / E2E", "Positive", "Negative", "Edge", "Ghi chú"]
COVER_COLS = {"Positive": 6, "Negative": 7, "Edge": 8}   # index trong tuple dòng RTM
FLOW_COL = 5
TC_ID_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)*_\d+")
NA_RE = re.compile(r"^\s*N/?A\b", re.I)
NA_REASON_RE = re.compile(r"^\s*N/?A\s*[—–:-]\s*\S.{3,}", re.I)
VAGUE = ["trạng thái mặc định", "đã có data", "user bình thường", "hoạt động bình thường", "như case trước"]
QNA_REF_RE = re.compile(r"Q&A\s*((?:#\d+\s*(?:,|và|/)?\s*)+)")


def read_qna(path):
    """{STT: Trạng thái} từ bảng markdown của file Q&A (cột đầu = STT, cột cuối = Trạng thái)."""
    status = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if line.lstrip().startswith("|") and cells and cells[0].isdigit():
                status[int(cells[0])] = cells[-1]
    return status


def lint_qna_refs(ws, name, qna):
    """(Q&A #n) phải có trong file Q&A; Expected 'Chờ confirm' mà câu đó đã chốt → phải cập nhật Expected."""
    header = [_cell_str(c.value) for c in ws[TC_HEADER_ROW]]
    id_i, exp_i = header.index("Test Case #"), header.index("Expected Result")
    errs = []
    for row in ws.iter_rows(min_row=TC_HEADER_ROW + 1, values_only=True):
        tc = _cell_str(row[id_i])
        if not tc:
            continue
        for i, val in enumerate(row):
            text = _cell_str(val)
            for m in QNA_REF_RE.finditer(text):
                for n in map(int, re.findall(r"#(\d+)", m.group(1))):
                    if n not in qna:
                        errs.append(f"[{name}] {tc}: tham chiếu Q&A #{n} không có trong file Q&A")
                    elif i == exp_i and "chờ confirm" in text.lower() and qna[n].lower().startswith("đã chốt"):
                        errs.append(f"[{name}] {tc}: Expected 'Chờ confirm' nhưng Q&A #{n} đã chốt — viết lại Expected theo câu trả lời")
    return errs


def lint_cases(ws, name):
    """Rule viết Pre-condition / Expected — kiểm theo thứ tự dòng trong sheet."""
    header = [_cell_str(c.value) for c in ws[TC_HEADER_ROW]]
    try:
        id_i, pre_i, exp_i = header.index("Test Case #"), header.index("Pre-condition"), header.index("Expected Result")
        step_i = header.index("Steps")
    except ValueError:
        return [f"[{name}] thiếu cột 'Pre-condition' / 'Expected Result'"]
    errs, prev = [], None
    for row in ws.iter_rows(min_row=TC_HEADER_ROW + 1, values_only=True):
        tc = _cell_str(row[id_i])
        if not tc:
            continue
        pre, exp = _cell_str(row[pre_i]), _cell_str(row[exp_i])
        E = lambda m: errs.append(f"[{name}] {tc}: {m}")
        if not exp:
            E("Expected Result trống")
        nums = lambda t: re.findall(r"^\s*(\d+)\.\s", t, re.M)
        s_nums, e_nums = nums(_cell_str(row[step_i])), nums(exp)
        if s_nums != e_nums:
            E(f"Expected phải có đúng 1 dòng cho mỗi step, cùng số (Steps {','.join(s_nums) or '—'} / Expected {','.join(e_nums) or '—'})")
        if not pre:
            E("Pre-condition trống (không có điều kiện thì ghi N/A)")
        elif pre.upper() != "N/A":
            m = re.match(r"Tiếp tục từ (\S+?):", pre)
            if m:
                if m.group(1) != prev:
                    E(f"'Tiếp tục từ {m.group(1)}' phải trỏ tới case ngay trước ({prev})")
            elif not (pre.startswith("Navigate tới màn ") or re.match(r"1\. ", pre)):
                E("Pre-condition phải mở đầu bằng '1. …' (dựng mới), 'Tiếp tục từ TC_xxx:' hoặc 'Navigate tới màn …'")
            for v in VAGUE:
                if v in pre.lower():
                    E(f"Pre-condition có câu mơ hồ '{v}' — thay bằng giá trị cụ thể")
        prev = tc
    return errs


def _cell_str(v):
    return "" if v is None else str(v).strip()


def read_testcases(ws):
    header = [_cell_str(c.value) for c in ws[TC_HEADER_ROW]]
    try:
        id_i, type_i = header.index("Test Case #"), header.index("Test Type")
    except ValueError:
        return None
    cases = {}
    for row in ws.iter_rows(min_row=TC_HEADER_ROW + 1, values_only=True):
        tc = _cell_str(row[id_i])
        if tc:
            cases[tc] = _cell_str(row[type_i])
    return cases


def read_rtm(ws):
    rows = []
    for row in ws.iter_rows(min_row=RTM_HEADER_ROW + 1, values_only=True):
        vals = [_cell_str(v) for v in (list(row) + [None] * len(RTM_COLS))[:len(RTM_COLS)]]
        if any(vals):
            rows.append(vals)
    return rows


def check_pair(tc_name, cases, rtm_name, rtm):
    errors = []
    E = lambda msg: errors.append(f"[{rtm_name}] {msg}")

    by_id, level = {}, {}
    for r in rtm:
        if not r[0]:
            E(f"Dòng thiếu ID: {r[3][:50]!r}")
            continue
        if r[0] in by_id:
            E(f"Trùng ID {r[0]}")
        by_id[r[0]] = r
        level[r[0]] = r[1]

    referenced = set()

    def refs(text, where):
        ids = TC_ID_RE.findall(text)
        for t in ids:
            if t not in cases:
                E(f"{where}: case {t} không có trong sheet '{tc_name}'")
        referenced.update(t for t in ids if t in cases)
        return [t for t in ids if t in cases]

    stats = {"Feature": [0, 0], "Flow": [0, 0], "Rule": [0, 0], "na": 0}
    for rid, r in by_id.items():
        lv, parent, flow_cell = r[1], r[2], r[FLOW_COL]
        if lv not in ("Feature", "Flow", "Rule"):
            E(f"{rid}: Cấp '{lv}' không hợp lệ")
            continue
        # cấu trúc cha-con
        if lv == "Feature" and parent:
            E(f"{rid}: Feature không được có 'Thuộc'")
        if lv == "Flow" and level.get(parent) != "Feature":
            E(f"{rid}: Flow phải thuộc 1 Feature có thật (đang là '{parent}')")
        if lv == "Rule" and level.get(parent) not in ("Flow", "Feature"):
            E(f"{rid}: Rule phải thuộc 1 Flow/Feature có thật (đang là '{parent}')")

        if lv in ("Feature", "Flow"):
            stats[lv][1] += 1
            if refs(flow_cell, f"{rid} (Flow / E2E)"):
                stats[lv][0] += 1
            else:
                E(f"{rid} ({lv}): thiếu case đi hết luồng ở cột 'Flow / E2E' — {r[3][:60]!r}")
            for col, i in COVER_COLS.items():
                refs(r[i], f"{rid} ({col})")
        else:
            refs(flow_cell, f"{rid} (Flow / E2E)")
            for col, i in COVER_COLS.items():
                stats["Rule"][1] += 1
                cell = r[i]
                if not cell:
                    E(f"{rid}: ô {col} trống — cần mã case hoặc 'N/A — <lý do>' — {r[3][:60]!r}")
                    continue
                if NA_RE.match(cell):
                    # rule lấy từ spec/design luôn phải có case Positive; chỉ dòng suy diễn từ checklist
                    # (vốn là case negative/edge) mới được để Positive là N/A
                    if (col == "Positive" and not r[4].lower().startswith("checklist")
                            and "không có hành vi positive" not in cell.lower()):
                        E(f"{rid}: ô Positive không được N/A — rule từ '{r[4] or 'spec'}' phải có ít nhất 1 case "
                          f"Positive (chỉ dòng Nguồn 'Checklist: ...' hoặc lý do 'không có hành vi positive' mới được N/A) — {r[3][:60]!r}")
                        continue
                    stats["na"] += 1
                    stats["Rule"][0] += 1
                    if not NA_REASON_RE.match(cell):
                        E(f"{rid}: ô {col} ghi N/A nhưng thiếu lý do cụ thể ('{cell}')")
                    continue
                ids = refs(cell, f"{rid} ({col})")
                if not ids:
                    E(f"{rid}: ô {col} không chứa mã case hợp lệ ('{cell}')")
                    continue
                stats["Rule"][0] += 1
                for t in ids:
                    if cases[t] != col:
                        E(f"{rid}: {t} nằm ở ô {col} nhưng Test Type của case là '{cases[t]}'")

    for fid, lv in level.items():
        if lv == "Feature" and not any(r[1] == "Flow" and r[2] == fid for r in by_id.values()):
            E(f"{fid}: Feature chưa có Flow nào")

    for t in sorted(set(cases) - referenced):
        E(f"Case {t} (sheet '{tc_name}') không được RTM tham chiếu — case mồ côi, gắn vào dòng RTM phù hợp hoặc xoá")

    return errors, stats


def main():
    if len(sys.argv) not in (2, 3):
        print(__doc__)
        sys.exit(2)
    qna = None
    if len(sys.argv) == 3:
        try:
            qna = read_qna(sys.argv[2])
        except OSError as e:
            print(f"Không đọc được file Q&A: {e}")
            sys.exit(2)
    try:
        wb = load_workbook(sys.argv[1], read_only=True)
    except Exception as e:  # file hỏng / sai đường dẫn
        print(f"Không đọc được file: {e}")
        sys.exit(2)

    pairs = []
    for ws in wb.worksheets:
        if _cell_str(ws.cell(row=1, column=1).value) == RTM_LINK_LABEL:
            pairs.append((_cell_str(ws.cell(row=1, column=2).value), ws.title))
    if not pairs:
        print("Không có sheet RTM nào (ô A1 = 'Sheet test case'). Build lại với tham số rtm=[...].")
        sys.exit(2)

    all_errors = []
    tc_sheets = {ws.title for ws in wb.worksheets} - {p[1] for p in pairs}
    for name in sorted(tc_sheets - {p[0] for p in pairs}):
        all_errors.append(f"[{name}] sheet test case chưa có RTM — chưa kiểm được bao phủ")
    for tc_name, rtm_name in pairs:
        if tc_name not in wb.sheetnames:
            all_errors.append(f"[{rtm_name}] trỏ tới sheet '{tc_name}' không tồn tại")
            continue
        cases = read_testcases(wb[tc_name])
        if cases is None:
            all_errors.append(f"[{tc_name}] không tìm thấy cột 'Test Case #'/'Test Type' ở dòng {TC_HEADER_ROW}")
            continue
        errors, st = check_pair(tc_name, cases, rtm_name, read_rtm(wb[rtm_name]))
        all_errors += errors + lint_cases(wb[tc_name], tc_name)
        if qna is not None:
            all_errors += lint_qna_refs(wb[tc_name], tc_name, qna)
        print(f"== {tc_name}: {len(cases)} case")
        print(f"   Feature có case E2E : {st['Feature'][0]}/{st['Feature'][1]}")
        print(f"   Flow có case E2E    : {st['Flow'][0]}/{st['Flow'][1]}")
        print(f"   Ô P/N/E của Rule    : {st['Rule'][0]}/{st['Rule'][1]} đã điền (trong đó {st['na']} ô N/A)")

    if all_errors:
        print(f"\n❌ {len(all_errors)} lỗi:")
        for e in all_errors:
            print("  - " + e)
        sys.exit(1)
    print("\n✅ RTM đạt: 0 lỗi.")


if __name__ == "__main__":
    main()
